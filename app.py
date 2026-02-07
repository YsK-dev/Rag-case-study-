import sys
import os
from pathlib import Path

# Add src to path BEFORE any relative imports
sys.path.insert(0, os.path.join(os.path.dirname(__file__), 'src'))

from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
import logging
import time
import json
from typing import Optional, List

from rag_engine import RAGEngine
from llm_client import OllamaClient

# Project paths
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
UPLOAD_DIR = DATA_DIR / "uploads"
STATIC_DIR = BASE_DIR / "static"

# Supported file types (aligned with RAGEngine)
SUPPORTED_EXTENSIONS = {
    ".pdf", ".txt", ".md", ".docx",
    ".xlsx", ".xlsm", ".csv",
    ".html", ".htm", ".pptx",
}

# Setup logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('rag_system.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# Initialize FastAPI app
app = FastAPI(
    title="RAG Q&A API",
    description="Retrieval-Augmented Generation API for document-based question answering",
    version="2.0.0"
)

# Add CORS middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Model mapping - frontend model names to Ollama models
MODEL_MAPPING = {
    "gemini-1.5-flash": "qwen3:1.7b",     # Fast model (smaller, quicker)
    "gemini-1.5-pro": "pielee/qwen3-4b-thinking-2507_q8:latest",  # Smart model (larger, better reasoning)
    "flash": "qwen3:1.7b",
    "pro": "pielee/qwen3-4b-thinking-2507_q8:latest"
}

# Initialize RAG components
print("Initializing RAG Engine...")
rag_engine = RAGEngine(
    db_path=str(BASE_DIR / "chroma_db_web"),
    collection_name="web_docs"
)
print("Initializing Ollama LLM...")
llm_client = OllamaClient(model="qwen3:1.7b")
print("Server ready!")

# Pydantic models
class ChatRequest(BaseModel):
    question: str
    top_k: int = 3
    model: Optional[str] = "gemini-1.5-flash"
    stream: Optional[bool] = False
    temperature: Optional[float] = 0.7  # Temperature for LLM generation (0.0-1.0)

class FeedbackRequest(BaseModel):
    message_id: str
    feedback: str  # "up" or "down"

class ReasoningTrace(BaseModel):
    query_rewrite: str = ""
    search_terms: List[str] = []
    retrieval_strategy: str = ""
    llm_thinking: Optional[str] = None

class SourceChunk(BaseModel):
    text: str
    source: str
    chunk_id: str = ""
    page: int = 0
    content_type: str = "text"
    file_path: str = ""
    char_count: int = 0
    token_count: int = 0
    char_count_raw: int = 0
    confidence: float = 0.5

class ChatResponse(BaseModel):
    answer: str
    sources: List[SourceChunk] = []
    metadata: dict = {}
    reasoning: Optional[ReasoningTrace] = None

# Feedback storage path
FEEDBACK_FILE = DATA_DIR / "feedback.json"

def load_feedback():
    if FEEDBACK_FILE.exists():
        with open(FEEDBACK_FILE, 'r') as f:
            return json.load(f)
    return {"entries": []}

def save_feedback(data):
    os.makedirs(FEEDBACK_FILE.parent, exist_ok=True)
    with open(FEEDBACK_FILE, 'w') as f:
        json.dump(data, f, indent=2)

# API Endpoints
@app.get("/")
async def root():
    """Serve the frontend if available, otherwise return API info."""
    index_path = STATIC_DIR / "index.html"
    if index_path.exists():
        return FileResponse(index_path)
    return {
        "status": "ok",
        "message": "RAG API is running",
        "docs": "/docs"
    }

@app.get("/api/health")
async def health_check():
    """Health check endpoint"""
    return {"status": "healthy", "model": llm_client.model}

@app.post("/api/feedback")
async def submit_feedback(request: FeedbackRequest):
    """Store user feedback for a message"""
    logger.info(f"Received feedback: {request.feedback} for message {request.message_id}")
    
    try:
        feedback_data = load_feedback()
        feedback_data["entries"].append({
            "message_id": request.message_id,
            "feedback": request.feedback,
            "timestamp": datetime.now().isoformat()
        })
        save_feedback(feedback_data)
        
        return {"status": "success", "message": "Feedback recorded"}
    except Exception as e:
        logger.error(f"Error saving feedback: {str(e)}")
        raise HTTPException(status_code=500, detail=str(e))

def generate_sse_stream(question: str, top_k: int, model_name: str, temperature: float = 0.7):
    """Generator for Server-Sent Events streaming response"""
    start_time = time.time()
    
    # Switch LLM model if needed
    llm_client.model = model_name
    
    try:
        # Analyze query for reasoning trace
        reasoning = rag_engine.analyze_query(question)
        
        # Retrieve relevant context
        retrieval_start = time.time()
        context_chunks, metadatas = rag_engine.retrieve(question, top_k=top_k)
        retrieval_time = time.time() - retrieval_start
        
        logger.info(f"Retrieved {len(context_chunks)} chunks in {retrieval_time:.2f}s")
        
        if not context_chunks:
            # No context found, use LLM directly with streaming
            generation_start = time.time()
            llm_thinking = ""
            for result in llm_client.chat_stream(question, temperature):
                if result.get('token'):  # Filter empty tokens
                    yield f"data: {json.dumps({'token': result['token']})}\n\n"
                if result.get('thinking'):
                    llm_thinking += result['thinking']
                    yield f"data: {json.dumps({'thinking': result['thinking']})}\n\n"
            generation_time = time.time() - generation_start
            
            total_time = time.time() - start_time
            
            # Send empty sources
            yield f"data: {json.dumps({'sources': []})}\n\n"
            
            # Send metadata
            metadata = {
                "model": model_name,
                "retrieval_time": round(retrieval_time, 3),
                "generation_time": round(generation_time, 3),
                "total_time": round(total_time, 3),
                "chunks_used": 0,
                "timestamp": datetime.now().isoformat()
            }
            yield f"data: {json.dumps({'metadata': metadata})}\n\n"
            
            # Send reasoning (include LLM thinking if present)
            if llm_thinking:
                reasoning['llm_thinking'] = llm_thinking
            yield f"data: {json.dumps({'reasoning': reasoning})}\n\n"
            
            yield "data: [DONE]\n\n"
            return
        
        # Stream answer generation with context
        generation_start = time.time()
        llm_thinking = ""
        for result in llm_client.generate_answer_stream(question, context_chunks, temperature):
            if result.get('token'):  # Filter empty tokens
                yield f"data: {json.dumps({'token': result['token']})}\n\n"
            if result.get('thinking'):
                llm_thinking += result['thinking']
                yield f"data: {json.dumps({'thinking': result['thinking']})}\n\n"
        generation_time = time.time() - generation_start
        
        total_time = time.time() - start_time
        
        # Format sources with metadata and confidence
        sources = []
        for i, (chunk, metadata) in enumerate(zip(context_chunks, metadatas)):
            source_file = metadata.get('source', 'Unknown')
            file_path = f"/api/documents/{source_file}" if source_file != 'Unknown' else ""
            
            sources.append({
                "text": chunk[:300] + "..." if len(chunk) > 300 else chunk,
                "source": source_file,
                "chunk_id": f"chunk_{i}",
                "page": metadata.get('page', 0),
                "content_type": metadata.get('type', 'text'),
                "file_path": file_path,
                "char_count": metadata.get('char_count', len(chunk)),
                "token_count": metadata.get('token_count', len(chunk) // 4),
                "confidence": metadata.get('confidence', 0.5)
            })
        
        # Send sources
        yield f"data: {json.dumps({'sources': sources})}\n\n"
        
        # Send metadata
        response_metadata = {
            "model": model_name,
            "retrieval_time": round(retrieval_time, 3),
            "generation_time": round(generation_time, 3),
            "total_time": round(total_time, 3),
            "chunks_used": len(context_chunks),
            "timestamp": datetime.now().isoformat()
        }
        yield f"data: {json.dumps({'metadata': response_metadata})}\n\n"
        
        # Send reasoning (include LLM thinking if present)
        if llm_thinking:
            reasoning['llm_thinking'] = llm_thinking
        yield f"data: {json.dumps({'reasoning': reasoning})}\n\n"
        
        yield "data: [DONE]\n\n"
        
    except Exception as e:
        logger.error(f"Streaming error: {str(e)}", exc_info=True)
        yield f"data: {json.dumps({'error': str(e)})}\n\n"
        yield "data: [DONE]\n\n"

@app.post("/api/chat")
async def chat(request: ChatRequest):
    """
    Answer questions using RAG with optional streaming
    """
    start_time = time.time()
    logger.info(f"Received question: {request.question}, stream={request.stream}, model={request.model}")
    
    # Get model name from mapping
    model_name = MODEL_MAPPING.get(request.model, "qwen3:1.7b")
    
    # Handle streaming request
    if request.stream:
        return StreamingResponse(
            generate_sse_stream(request.question, request.top_k, model_name, request.temperature),
            media_type="text/event-stream",
            headers={
                "Cache-Control": "no-cache",
                "Connection": "keep-alive",
                "X-Accel-Buffering": "no"
            }
        )
    
    # Non-streaming response
    try:
        # Switch LLM model if needed
        llm_client.model = model_name
        
        # Analyze query for reasoning trace
        reasoning = rag_engine.analyze_query(request.question)
        
        # Retrieve relevant context
        retrieval_start = time.time()
        context_chunks, metadatas = rag_engine.retrieve(request.question, top_k=request.top_k)
        retrieval_time = time.time() - retrieval_start
        
        logger.info(f"Retrieved {len(context_chunks)} chunks in {retrieval_time:.2f}s")
        
        if not context_chunks:
            # No context found, use LLM directly
            logger.warning("No relevant context found, using LLM without RAG")
            generation_start = time.time()
            answer = llm_client.chat(request.question, request.temperature)
            generation_time = time.time() - generation_start
            
            total_time = time.time() - start_time
            logger.info(f"Response generated in {generation_time:.2f}s (total: {total_time:.2f}s)")
            
            return ChatResponse(
                answer=answer,
                sources=[],
                metadata={
                    "model": model_name,
                    "retrieval_time": round(retrieval_time, 3),
                    "generation_time": round(generation_time, 3),
                    "total_time": round(total_time, 3),
                    "chunks_used": 0,
                    "timestamp": datetime.now().isoformat()
                },
                reasoning=ReasoningTrace(**reasoning, llm_thinking=llm_client.last_thinking)
            )
        
        # Generate answer with context
        generation_start = time.time()
        answer, _thinking = llm_client.generate_answer(request.question, context_chunks, request.temperature)
        generation_time = time.time() - generation_start
        
        total_time = time.time() - start_time
        logger.info(f"Answer generated in {generation_time:.2f}s (total: {total_time:.2f}s)")
        
        # Format sources with metadata and confidence
        sources = []
        for i, (chunk, metadata) in enumerate(zip(context_chunks, metadatas)):
            # Get file path for the uploaded document
            source_file = metadata.get('source', 'Unknown')
            file_path = f"/api/documents/{source_file}" if source_file != 'Unknown' else ""
            
            sources.append(SourceChunk(
                text=chunk[:300] + "..." if len(chunk) > 300 else chunk,
                source=source_file,
                chunk_id=f"chunk_{i}",
                page=metadata.get('page', 0),
                content_type=metadata.get('type', 'text'),
                file_path=file_path,
                char_count=metadata.get('char_count', len(chunk)),
                token_count=metadata.get('token_count', len(chunk) // 4),
                char_count_raw=metadata.get('char_count_raw', len(chunk)),
                confidence=metadata.get('confidence', 0.5)
            ))
        
        return ChatResponse(
            answer=answer,
            sources=sources,
            metadata={
                "model": model_name,
                "retrieval_time": round(retrieval_time, 3),
                "generation_time": round(generation_time, 3),
                "total_time": round(total_time, 3),
                "chunks_used": len(context_chunks),
                "timestamp": datetime.now().isoformat()
            },
            reasoning=ReasoningTrace(**reasoning, llm_thinking=_thinking)
        )
        
    except Exception as e:
        logger.error(f"Error processing question: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))

@app.get("/api/pdf/{filename}")
async def serve_pdf(filename: str):
    """Legacy PDF endpoint — redirects to the generic documents endpoint."""
    from fastapi.responses import RedirectResponse
    return RedirectResponse(url=f"/api/documents/{filename}")

@app.get("/api/documents/{filename}")
async def get_document(filename: str):
    """Serve uploaded documents for preview (PDF, DOCX, XLSX, CSV, MD, TXT, HTML, PPTX)."""
    safe_name = Path(filename).name
    ext = Path(safe_name).suffix.lower()

    mime_map = {
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
        ".csv": "text/csv",
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".html": "text/html",
        ".htm": "text/html",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    }

    if ext not in mime_map:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type: {ext}"
        )

    file_path = UPLOAD_DIR / safe_name
    if not file_path.exists():
        file_path = DATA_DIR / safe_name
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Document not found")

    return FileResponse(
        file_path,
        media_type=mime_map[ext],
        filename=safe_name
    )

@app.post("/api/upload")
async def upload_document(file: UploadFile = File(...)):
    """
    Upload and ingest a document (PDF, DOCX, XLSX, CSV, TXT, MD, HTML, PPTX)
    """
    start_time = time.time()
    logger.info(f"Uploading file: {file.filename}")
    
    try:
        # Validate file type
        safe_name = Path(file.filename).name
        file_ext = Path(safe_name).suffix.lower()
        if file_ext not in SUPPORTED_EXTENSIONS:
            logger.warning(f"Invalid file type: {file.filename}")
            allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported file type. Allowed: {allowed}"
            )
        
        # Save uploaded file
        os.makedirs(UPLOAD_DIR, exist_ok=True)
        
        file_path = UPLOAD_DIR / safe_name
        file_size = 0
        
        with open(file_path, "wb") as buffer:
            content = await file.read()
            file_size = len(content)
            buffer.write(content)
        
        logger.info(f"File saved to: {file_path}")
        
        # Ingest the document
        ingest_start = time.time()
        result = rag_engine.ingest_file(file_path)
        ingest_time = time.time() - ingest_start
        
        total_time = time.time() - start_time
        logger.info(f"Document ingested in {ingest_time:.2f}s (total: {total_time:.2f}s)")
        
        # Parse chunks count from result message
        chunks_created = 0
        if "chunks" in result.lower():
            try:
                chunks_created = int(result.split()[2])  # "Successfully processed X chunks..."
            except (ValueError, IndexError):
                chunks_created = 0
        
        return {
            "status": "success",
            "message": result,
            "filename": safe_name,
            "file_size": file_size,
            "chunks_created": chunks_created,
            "processing_time": round(total_time, 2)
        }
        
    except Exception as e:
        logger.error(f"Error uploading file: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/api/documents/{filename}/preview")
async def preview_document(filename: str):
    """Extract text/table content from non-PDF documents for in-browser preview."""
    safe_name = Path(filename).name
    ext = Path(safe_name).suffix.lower()

    previewable = {".docx", ".xlsx", ".xlsm", ".csv", ".pptx", ".txt", ".md", ".html", ".htm"}
    if ext not in previewable:
        raise HTTPException(status_code=400, detail=f"Preview not supported for {ext}")

    file_path = UPLOAD_DIR / safe_name
    if not file_path.exists():
        file_path = DATA_DIR / safe_name
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Document not found")

    try:
        if ext in (".txt", ".md"):
            with open(file_path, "r", encoding="utf-8", errors="replace") as fh:
                return {"type": "text", "content": fh.read(), "filename": safe_name}

        if ext in (".html", ".htm"):
            with open(file_path, "r", encoding="utf-8", errors="replace") as fh:
                return {"type": "html", "content": fh.read(), "filename": safe_name}

        if ext == ".csv":
            import csv as csv_mod
            rows = []
            with open(file_path, "r", newline="", encoding="utf-8", errors="replace") as fh:
                reader = csv_mod.reader(fh)
                for row in reader:
                    rows.append(row)
            return {"type": "table", "sheets": [{"name": safe_name, "rows": rows}], "filename": safe_name}

        if ext in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook as _load_wb
            wb = _load_wb(str(file_path), data_only=True, read_only=True)
            sheets = []
            for sheet in wb.worksheets:
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(["" if c is None else str(c) for c in row])
                sheets.append({"name": sheet.title, "rows": rows})
            return {"type": "table", "sheets": sheets, "filename": safe_name}

        if ext == ".docx":
            from docx import Document as _Document
            doc = _Document(str(file_path))
            paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
            tables = []
            for table in doc.tables:
                tbl_rows = []
                for row in table.rows:
                    tbl_rows.append([cell.text.strip() for cell in row.cells])
                tables.append(tbl_rows)
            return {
                "type": "docx",
                "paragraphs": paragraphs,
                "tables": tables,
                "filename": safe_name,
            }

        if ext == ".pptx":
            from pptx import Presentation as _Presentation
            prs = _Presentation(str(file_path))
            slides = []
            for slide in prs.slides:
                texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        texts.append(shape.text)
                slides.append("\n".join(texts))
            return {"type": "pptx", "slides": slides, "filename": safe_name}

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Preview extraction failed: {str(e)}") from e

    raise HTTPException(status_code=400, detail=f"Preview not supported for {ext}")

# Mount static files
if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
else:
    logger.warning("Static directory not found; skipping static file mount.")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
