"""RAG Engine module for document processing and retrieval."""

import csv
import hashlib
import json as json_mod
import math
import os
import re
import uuid
from collections import Counter
from typing import Dict, List, Optional, Tuple

import chromadb
import fitz  # PyMuPDF
import numpy as np
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from sentence_transformers import CrossEncoder, SentenceTransformer

try:
    import faiss

    FAISS_AVAILABLE = True
except ImportError:
    FAISS_AVAILABLE = False

# --------------- BM25 (keyword search) ---------------
try:
    from rank_bm25 import BM25Okapi  # type: ignore

    BM25_AVAILABLE = True
except ImportError:
    BM25_AVAILABLE = False


class RAGEngine:
    """Enhanced RAG Engine with PDF processing and two-stage retrieval."""

    def __init__(
        self,
        db_path: str = "./chroma_db",
        collection_name: str = "docs_collection",
        backend: str = "chroma",
        embedding_model: Optional[SentenceTransformer] = None,
        reranker: Optional[CrossEncoder] = None,
    ):
        """Initialize the Enhanced RAG Engine.

        Args:
            db_path: Path to persistent storage.
            collection_name: Name of the ChromaDB collection (chroma only).
            backend: Vector store backend – ``"chroma"`` or ``"faiss"``.
            embedding_model: Pre-loaded SentenceTransformer (avoids reload).
            reranker: Pre-loaded CrossEncoder reranker (avoids reload).
        """
        self.backend = backend

        # --- Shared models (load once, reuse across backends) ---
        if embedding_model is not None:
            self.embedding_model = embedding_model
        else:
            print("Loading embedding model... (this happens only once)")
            self.embedding_model = SentenceTransformer("all-MiniLM-L6-v2")

        if reranker is not None:
            self.reranker = reranker
        else:
            print("Loading reranker model...")
            self.reranker = CrossEncoder(
                "cross-encoder/ms-marco-MiniLM-L-6-v2"
            )

        # --- Backend-specific initialisation ---
        if backend == "chroma":
            self.chroma_client = chromadb.PersistentClient(path=db_path)
            self.collection = self.chroma_client.get_or_create_collection(
                name=collection_name
            )
        elif backend == "faiss":
            if not FAISS_AVAILABLE:
                raise ImportError(
                    "faiss-cpu is required for the FAISS backend. "
                    "Install with: pip install faiss-cpu"
                )
            self._init_faiss(db_path)
        else:
            raise ValueError(
                f"Unknown backend: {backend}. Use 'chroma' or 'faiss'."
            )

        # --- BM25 cache (avoids rebuilding index on every query) ---
        self._bm25_cache: Optional[Dict] = None

        print(f"Models loaded. Backend: {backend}")

    # ------------------------------------------------------------------
    # FAISS helpers
    # ------------------------------------------------------------------

    def _init_faiss(self, db_path: str) -> None:
        """Create or load a FAISS index with sidecar metadata."""
        self.faiss_db_path = db_path
        os.makedirs(db_path, exist_ok=True)

        self.embedding_dim = 384  # all-MiniLM-L6-v2 output dim

        index_path = os.path.join(db_path, "faiss.index")
        meta_path = os.path.join(db_path, "faiss_meta.json")

        if os.path.exists(index_path) and os.path.exists(meta_path):
            self.faiss_index = faiss.read_index(index_path)
            with open(meta_path, "r", encoding="utf-8") as fh:
                sidecar = json_mod.load(fh)
            self.faiss_documents: List[str] = sidecar["documents"]
            self.faiss_metadatas: List[Dict] = sidecar["metadatas"]
            self.faiss_ids: List[str] = sidecar["ids"]
            print(
                f"Loaded FAISS index with "
                f"{self.faiss_index.ntotal} vectors"
            )
        else:
            self.faiss_index = faiss.IndexFlatL2(self.embedding_dim)
            self.faiss_documents = []
            self.faiss_metadatas = []
            self.faiss_ids = []
            print("Created new FAISS index (IndexFlatL2)")

    def _save_faiss(self) -> None:
        """Persist the FAISS index and sidecar metadata to disk."""
        index_path = os.path.join(self.faiss_db_path, "faiss.index")
        meta_path = os.path.join(self.faiss_db_path, "faiss_meta.json")

        faiss.write_index(self.faiss_index, index_path)
        with open(meta_path, "w", encoding="utf-8") as fh:
            json_mod.dump(
                {
                    "documents": self.faiss_documents,
                    "metadatas": self.faiss_metadatas,
                    "ids": self.faiss_ids,
                },
                fh,
            )

    def _faiss_add(
        self,
        texts: List[str],
        embeddings: List[List[float]],
        metadatas: List[Dict],
        ids: List[str],
    ) -> None:
        """Add vectors + metadata to the FAISS store."""
        vectors = np.array(embeddings, dtype=np.float32)
        self.faiss_index.add(vectors)
        self.faiss_documents.extend(texts)
        self.faiss_metadatas.extend(metadatas)
        self.faiss_ids.extend(ids)
        self._save_faiss()

    def _faiss_query(
        self, query_embedding: List[float], n_results: int
    ) -> Dict:
        """Query the FAISS index and return chroma-compatible dict."""
        if self.faiss_index.ntotal == 0:
            return {
                "documents": [[]],
                "metadatas": [[]],
                "distances": [[]],
            }

        query_vec = np.array(
            [query_embedding], dtype=np.float32
        )
        n_results = min(n_results, self.faiss_index.ntotal)

        distances, indices = self.faiss_index.search(
            query_vec, n_results
        )

        docs: List[str] = []
        metas: List[Dict] = []
        dists: List[float] = []

        for i, idx in enumerate(indices[0]):
            if idx < 0:  # FAISS returns -1 for unfilled slots
                continue
            docs.append(self.faiss_documents[idx])
            metas.append(self.faiss_metadatas[idx])
            dists.append(float(distances[0][i]))

        return {
            "documents": [docs],
            "metadatas": [metas],
            "distances": [dists],
        }

    def get_stats(self) -> Dict:
        """Return basic stats about the current vector store."""
        if self.backend == "faiss":
            return {
                "backend": "faiss",
                "total_vectors": self.faiss_index.ntotal,
                "index_type": type(self.faiss_index).__name__,
                "dimension": self.embedding_dim,
            }
        # chroma
        count = self.collection.count()
        return {
            "backend": "chroma",
            "total_vectors": count,
            "collection": self.collection.name,
        }

    # ------------------------------------------------------------------
    # Content-type detection
    # ------------------------------------------------------------------

    def _detect_content_type(self, text: str) -> str:
        """Detect if text contains equations, code, or tables.

        Args:
            text: The text content to analyze.

        Returns:
            Content type: 'equation', 'code', 'table', or 'text'.
        """
        if self._is_equation(text):
            return "equation"
        if self._is_code(text):
            return "code"
        if self._is_table(text):
            return "table"
        return "text"

    @staticmethod
    def _is_equation(text: str) -> bool:
        """Check whether text looks like a math equation."""
        math_indicators = [
            r"\$\$.*?\$\$",
            r"\$.*?\$",
            r"\\text\{",
            r"\\frac\{",
            r"\\sum",
            r"\\int",
            r"\\left",
            r"\\right",
            r"[\u2211\u222b\u2202\u2207\u2248\u2260\u2264\u2265\u00b1\u00d7\u00f7]",
        ]
        math_count = sum(
            1 for p in math_indicators if re.search(p, text)
        )
        return math_count >= 2

    @staticmethod
    def _is_code(text: str) -> bool:
        """Check whether text looks like source code."""
        if re.search(r"^\s*(```|~~~)", text, re.MULTILINE):
            return True
        if re.search(r"</?(code|pre)\b", text, re.IGNORECASE):
            return True

        code_patterns = [
            r"^\s{4,}\w+",
            r"\bdef\s+\w+\s*\(",
            r"\bclass\s+\w+",
            r"\bimport\s+\w+",
            r"\breturn\s+",
            r"^\w+\s*=\s*\w+\(",
            r"[{};]\s*$",
        ]
        code_count = sum(
            1
            for p in code_patterns
            if re.search(p, text, re.MULTILINE)
        )
        return code_count >= 2

    @staticmethod
    def _is_table(text: str) -> bool:
        """Check whether text looks like a table."""
        if re.search(r"^\s*\|.*\|\s*$", text, re.MULTILINE):
            if re.search(
                r"^\s*\|?\s*:?-{3,}\s*\|", text, re.MULTILINE
            ):
                return True
            if text.count("|") >= 4:
                return True

        lines = [
            line for line in text.splitlines() if line.strip()
        ]
        comma_lines = [
            line for line in lines if line.count(",") >= 2
        ]
        if len(comma_lines) >= 2:
            return True

        return text.count("\t") > 3

    # ------------------------------------------------------------------
    # File-reading helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _read_text_file(file_path: str) -> str:
        """Read a text-based file with encoding fallbacks.

        Args:
            file_path: Path to the text file.

        Returns:
            The file content as a string.
        """
        encodings = ["utf-8", "utf-8-sig", "latin-1"]
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as fh:
                    return fh.read()
            except UnicodeDecodeError:
                continue
        with open(
            file_path, "r", encoding="utf-8", errors="replace"
        ) as fh:
            return fh.read()

    def _extract_text_to_chunks(
        self,
        text: str,
        source: str,
        page: int = 1,
        chunk_size: int = 500,
        chunk_overlap: int = 100,
    ) -> List[Dict]:
        """Convert raw text into chunk dicts with metadata using a sliding window.

        Instead of splitting on blank lines with no overlap, this uses
        a sentence-aware sliding window so context at chunk boundaries
        is preserved.

        Args:
            text: Raw text to process.
            source: Source file name.
            page: Page number for the metadata.
            chunk_size: Target character count per chunk.
            chunk_overlap: Number of overlapping characters between
                consecutive chunks.

        Returns:
            List of chunk dictionaries.
        """
        if not text or not text.strip():
            return []

        # Split into sentences (keep delimiters attached)
        sentences = re.split(r'(?<=[.!?;\n])\s+', text.strip())
        sentences = [s.strip() for s in sentences if s.strip()]

        if not sentences:
            return []

        chunks: List[Dict] = []
        current_chunk: List[str] = []
        current_len = 0

        for sentence in sentences:
            sen_len = len(sentence)

            if current_len + sen_len > chunk_size and current_chunk:
                chunk_text = " ".join(current_chunk).strip()
                if len(chunk_text) >= 50:
                    content_type = self._detect_content_type(chunk_text)
                    chunks.append(
                        {
                            "text": chunk_text,
                            "page": page,
                            "type": content_type,
                            "source": source,
                        }
                    )

                # Keep sentences that fit within the overlap window
                overlap_chars = 0
                overlap_sentences: List[str] = []
                for s in reversed(current_chunk):
                    overlap_chars += len(s)
                    if overlap_chars > chunk_overlap:
                        break
                    overlap_sentences.insert(0, s)

                current_chunk = overlap_sentences
                current_len = sum(len(s) for s in current_chunk)

            current_chunk.append(sentence)
            current_len += sen_len

        # Flush the remaining buffer
        if current_chunk:
            chunk_text = " ".join(current_chunk).strip()
            if len(chunk_text) >= 50:
                content_type = self._detect_content_type(chunk_text)
                chunks.append(
                    {
                        "text": chunk_text,
                        "page": page,
                        "type": content_type,
                        "source": source,
                    }
                )

        return chunks

    @staticmethod
    def _split_table_rows(
        rows: List[List[str]], max_rows: int = 50
    ) -> List[List[List[str]]]:
        """Split large tables into smaller chunks.

        Uses the first row as a repeated header when splitting.

        Args:
            rows: Table rows.
            max_rows: Maximum rows per chunk.

        Returns:
            List of table chunks.
        """
        if not rows:
            return []
        if len(rows) <= max_rows:
            return [rows]

        header = rows[0]
        body = rows[1:]
        table_chunks: List[List[List[str]]] = []
        chunk_body_size = max_rows - 1

        for i in range(0, len(body), chunk_body_size):
            table_chunks.append(
                [header] + body[i : i + chunk_body_size]
            )

        return table_chunks

    # ------------------------------------------------------------------
    # Per-format extractors
    # ------------------------------------------------------------------

    def _extract_with_layout(self, file_path: str) -> List[Dict]:
        """Extract text from PDF with layout preservation.

        Args:
            file_path: Path to the PDF file.

        Returns:
            List of chunk dicts with metadata.
        """
        chunks: List[Dict] = []
        source = os.path.basename(file_path)

        with fitz.open(file_path) as doc:
            for page_index, page in enumerate(doc):
                page_num = page_index + 1
                text = page.get_text("text", sort=True)
                chunks.extend(
                    self._extract_text_to_chunks(
                        text, source, page_num
                    )
                )

                tables = self._extract_tables_pymupdf(page)
                for table in tables:
                    for trows in self._split_table_rows(table):
                        if trows:
                            ttext = self._format_table(trows)
                            chunks.append(
                                {
                                    "text": f"[TABLE]\n{ttext}",
                                    "page": page_num,
                                    "type": "table",
                                    "source": source,
                                }
                            )

        return chunks

    @staticmethod
    def _extract_tables_pymupdf(page) -> List[List[List[str]]]:
        """Best-effort table extraction with PyMuPDF.

        Args:
            page: A PyMuPDF page object.

        Returns:
            List of tables, each table being a list of rows.
        """
        if not hasattr(page, "find_tables"):
            return []

        try:
            table_finder = page.find_tables()
        except Exception:  # pylint: disable=broad-except
            return []

        if not table_finder or not getattr(
            table_finder, "tables", None
        ):
            return []

        tables: List[List[List[str]]] = []
        for table in table_finder.tables:
            try:
                table_rows = table.extract()
            except Exception:  # pylint: disable=broad-except
                continue
            if table_rows:
                tables.append(table_rows)

        return tables

    def _extract_text_file(self, file_path: str) -> List[Dict]:
        """Extract chunks from a plain-text or Markdown file."""
        text = self._read_text_file(file_path)
        source = os.path.basename(file_path)
        return self._extract_text_to_chunks(text, source, page=1)

    def _extract_docx(self, file_path: str) -> List[Dict]:
        """Extract chunks from a DOCX file."""
        doc = Document(file_path)
        source = os.path.basename(file_path)

        paragraphs = [
            p.text.strip()
            for p in doc.paragraphs
            if p.text and p.text.strip()
        ]
        text = "\n\n".join(paragraphs)
        chunks = self._extract_text_to_chunks(
            text, source, page=1
        )

        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [
                    cell.text.strip() if cell.text else ""
                    for cell in row.cells
                ]
                rows.append(cells)

            for trows in self._split_table_rows(rows):
                if trows:
                    ttext = self._format_table(trows)
                    chunks.append(
                        {
                            "text": f"[TABLE]\n{ttext}",
                            "page": 1,
                            "type": "table",
                            "source": source,
                        }
                    )

        return chunks

    def _extract_xlsx(self, file_path: str) -> List[Dict]:
        """Extract chunks from an Excel file."""
        workbook = load_workbook(
            file_path, data_only=True, read_only=True
        )
        source = os.path.basename(file_path)
        chunks: List[Dict] = []

        for sheet_idx, sheet in enumerate(workbook.worksheets, 1):
            rows: List[List[str]] = []
            for row in sheet.iter_rows(values_only=True):
                clean = [
                    "" if c is None else str(c).strip()
                    for c in row
                ]
                if any(cell for cell in clean):
                    rows.append(clean)

            for trows in self._split_table_rows(rows):
                if trows:
                    ttext = self._format_table(trows)
                    chunks.append(
                        {
                            "text": (
                                f"[TABLE]\nSheet: {sheet.title}"
                                f"\n{ttext}"
                            ),
                            "page": sheet_idx,
                            "type": "table",
                            "source": source,
                        }
                    )

        return chunks

    def _extract_csv(self, file_path: str) -> List[Dict]:
        """Extract chunks from a CSV file."""
        encodings = ["utf-8-sig", "utf-8", "latin-1"]
        rows: List[List[str]] = []

        for enc in encodings:
            try:
                with open(
                    file_path, "r", newline="", encoding=enc
                ) as fh:
                    reader = csv.reader(fh)
                    for row in reader:
                        clean = [c.strip() for c in row]
                        if any(clean):
                            rows.append(clean)
                break
            except UnicodeDecodeError:
                continue

        if not rows:
            with open(
                file_path,
                "r",
                newline="",
                encoding="utf-8",
                errors="replace",
            ) as fh:
                reader = csv.reader(fh)
                for row in reader:
                    clean = [c.strip() for c in row]
                    if any(clean):
                        rows.append(clean)

        source = os.path.basename(file_path)
        chunks: List[Dict] = []

        for trows in self._split_table_rows(rows):
            if trows:
                ttext = self._format_table(trows)
                chunks.append(
                    {
                        "text": f"[TABLE]\n{ttext}",
                        "page": 1,
                        "type": "table",
                        "source": source,
                    }
                )

        return chunks

    def _extract_html(self, file_path: str) -> List[Dict]:
        """Extract chunks from an HTML file."""
        html = self._read_text_file(file_path)
        soup = BeautifulSoup(html, "html.parser")

        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()

        text = soup.get_text(separator="\n")
        text = re.sub(r"\n{2,}", "\n\n", text).strip()

        source = os.path.basename(file_path)
        return self._extract_text_to_chunks(
            text, source, page=1
        )

    def _extract_pptx(self, file_path: str) -> List[Dict]:
        """Extract chunks from a PowerPoint file."""
        prs = Presentation(file_path)
        source = os.path.basename(file_path)
        chunks: List[Dict] = []

        for slide_idx, slide in enumerate(prs.slides, 1):
            text_parts: List[str] = []

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text_parts.append(shape.text)

                if getattr(shape, "has_table", False):
                    rows = []
                    for row in shape.table.rows:
                        cells = [
                            cell.text.strip()
                            if cell.text
                            else ""
                            for cell in row.cells
                        ]
                        rows.append(cells)

                    for trows in self._split_table_rows(rows):
                        if trows:
                            ttext = self._format_table(trows)
                            chunks.append(
                                {
                                    "text": f"[TABLE]\n{ttext}",
                                    "page": slide_idx,
                                    "type": "table",
                                    "source": source,
                                }
                            )

            slide_text = "\n\n".join(
                part.strip()
                for part in text_parts
                if part and part.strip()
            )
            chunks.extend(
                self._extract_text_to_chunks(
                    slide_text, source, page=slide_idx
                )
            )

        return chunks

    # ------------------------------------------------------------------
    # Formatting / chunking helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _format_table(table: List[List]) -> str:
        """Format table data as a markdown table string.

        Args:
            table: Table data as a list of rows.

        Returns:
            Markdown-formatted table string.
        """
        if not table:
            return ""

        lines = []
        for row in table:
            clean = [str(c) if c else "" for c in row]
            lines.append("| " + " | ".join(clean) + " |")

        if len(lines) > 1:
            num_cols = len(table[0])
            sep = "| " + " | ".join(["---"] * num_cols) + " |"
            lines.insert(1, sep)

        return "\n".join(lines)

    def _smart_chunk(
        self, chunks: List[Dict], max_chunk_size: int = 1000
    ) -> List[Dict]:
        """Smart chunking that respects content boundaries.

        Args:
            chunks: List of raw extracted chunks.
            max_chunk_size: Maximum character count per chunk.

        Returns:
            List of processed chunks with metadata.
        """
        smart_chunks: List[Dict] = []
        cur: Dict = {"text": "", "metadata": {}}

        for chunk in chunks:
            chunk_text = chunk["text"]

            if chunk["type"] in ["table", "code", "equation"]:
                if cur["text"]:
                    smart_chunks.append(cur)
                    cur = {"text": "", "metadata": {}}

                smart_chunks.append(
                    {
                        "text": chunk_text,
                        "metadata": self._build_meta(
                            chunk, chunk_text
                        ),
                    }
                )
                continue

            combined = len(cur["text"]) + len(chunk_text)
            if combined < max_chunk_size:
                if cur["text"]:
                    cur["text"] += "\n\n" + chunk_text
                else:
                    cur["text"] = chunk_text
                cur["metadata"] = self._build_meta(
                    chunk, cur["text"]
                )
            else:
                if cur["text"]:
                    smart_chunks.append(cur)
                cur = {
                    "text": chunk_text,
                    "metadata": self._build_meta(
                        chunk, chunk_text
                    ),
                }

        if cur["text"]:
            smart_chunks.append(cur)

        return smart_chunks

    @staticmethod
    def _build_meta(chunk: Dict, text: str) -> Dict:
        """Build metadata dict for a chunk.

        Args:
            chunk: Original chunk dict with page, type, source.
            text: The actual text content.

        Returns:
            Metadata dictionary.
        """
        return {
            "page": chunk["page"],
            "type": chunk.get("type", "text"),
            "source": chunk["source"],
            "char_count": len(text),
            "token_count": len(text) // 4,
            "char_count_raw": len(text),
        }

    # ------------------------------------------------------------------
    # Deduplication
    # ------------------------------------------------------------------

    @staticmethod
    def _normalize_for_dedup(text: str) -> str:
        """Normalize text for deduplication comparison.

        Lowercases, collapses whitespace, and strips
        punctuation so that near-identical repeated phrases
        (headers, footers, disclaimers) match reliably.

        Args:
            text: Raw chunk text.

        Returns:
            Normalized string used only for comparison.
        """
        t = text.lower()
        t = re.sub(r"[^\w\s]", "", t)   # remove punctuation
        t = re.sub(r"\s+", " ", t).strip()
        return t

    @staticmethod
    def _fingerprint(text: str) -> str:
        """Return a short hash of the normalized text."""
        return hashlib.md5(text.encode("utf-8")).hexdigest()

    def _deduplicate_chunks(
        self,
        chunks: List[Dict],
        max_repeat: int = 2,
        similarity_threshold: float = 0.95,
    ) -> Tuple[List[Dict], int]:
        """Remove near-duplicate chunks, keeping at most *max_repeat* copies.

        Two-layer check:
        1. **Exact match** (MD5 of normalised text) — very fast.
        2. **Near-duplicate** (short text overlap ratio) — catches
           minor wording differences like page numbers embedded in
           repeated headers.

        Args:
            chunks: Smart-chunked list ready for ingestion.
            max_repeat: How many times the same content is allowed
                        (default ``2``).  Set to ``1`` for strict
                        dedup, ``3`` to be more lenient.
            similarity_threshold: Jaccard word-overlap ratio above
                                  which two chunks are considered
                                  near-duplicates (0.0–1.0).

        Returns:
            Tuple of (deduplicated chunks, number removed).
        """
        fingerprint_counts: Counter = Counter()
        seen_word_sets: List[Tuple[set, str]] = []
        kept: List[Dict] = []
        removed = 0

        for chunk in chunks:
            norm = self._normalize_for_dedup(chunk["text"])

            # --- Layer 1: exact fingerprint ---
            fp = self._fingerprint(norm)
            if fingerprint_counts[fp] >= max_repeat:
                removed += 1
                continue

            # --- Layer 2: near-duplicate (word-level Jaccard) ---
            words = set(norm.split())
            is_near_dup = False

            if len(words) <= 60:  # only check short chunks
                for seen_words, seen_fp in seen_word_sets:
                    if not words or not seen_words:
                        continue
                    intersection = words & seen_words
                    union = words | seen_words
                    jaccard = len(intersection) / len(union)
                    if jaccard >= similarity_threshold:
                        # treat as same content
                        if fingerprint_counts[seen_fp] >= max_repeat:
                            is_near_dup = True
                            break
                        # count against the earlier fingerprint
                        fingerprint_counts[seen_fp] += 1
                        fp = seen_fp  # unify identity
                        break

            if is_near_dup:
                removed += 1
                continue

            fingerprint_counts[fp] += 1
            seen_word_sets.append((words, fp))
            kept.append(chunk)

        return kept, removed

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def ingest_file(self, file_path: str) -> str:
        """Read a supported file and store chunks in ChromaDB.

        Args:
            file_path: Path to the file to ingest.

        Returns:
            Success message with the number of chunks processed.

        Raises:
            FileNotFoundError: If the file does not exist.
            ValueError: If the file type is not supported.
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(
                f"File not found: {file_path}"
            )

        ext = os.path.splitext(file_path)[1].lower()
        extractors = {
            ".pdf": self._extract_with_layout,
            ".txt": self._extract_text_file,
            ".md": self._extract_text_file,
            ".docx": self._extract_docx,
            ".xlsx": self._extract_xlsx,
            ".xlsm": self._extract_xlsx,
            ".csv": self._extract_csv,
            ".html": self._extract_html,
            ".htm": self._extract_html,
            ".pptx": self._extract_pptx,
        }

        extractor = extractors.get(ext)
        if extractor is None:
            raise ValueError(f"Unsupported file type: {ext}")

        raw_chunks = extractor(file_path)
        processed = self._smart_chunk(raw_chunks)

        if not processed:
            basename = os.path.basename(file_path)
            return f"No extractable content found in {basename}"

        # Deduplicate: cap repeated phrases at max_repeat copies
        processed, dupes_removed = self._deduplicate_chunks(
            processed, max_repeat=2
        )

        if not processed:
            basename = os.path.basename(file_path)
            return (
                f"No unique content found in {basename}"
                f" ({dupes_removed} duplicate chunks removed)"
            )

        texts = [c["text"] for c in processed]
        metadatas = [c["metadata"] for c in processed]
        batch_id = uuid.uuid4().hex[:8]
        ids = [
            f"{batch_id}_chunk_{i}" for i in range(len(texts))
        ]

        embeddings = self.embedding_model.encode(texts).tolist()

        if self.backend == "faiss":
            self._faiss_add(texts, embeddings, metadatas, ids)
        else:
            self.collection.add(
                documents=texts,
                embeddings=embeddings,
                metadatas=metadatas,
                ids=ids,
            )

        # Invalidate BM25 cache (corpus changed)
        self._bm25_cache = None

        basename = os.path.basename(file_path)
        dedup_msg = (
            f" ({dupes_removed} duplicate chunks removed)"
            if dupes_removed > 0
            else ""
        )
        return (
            f"Successfully processed {len(texts)} chunks"
            f" from {basename}{dedup_msg}"
        )

    def retrieve(
        self,
        query: str,
        top_k: int = 3,
        confidence_threshold: float = 0.3,
        source_filter: Optional[str] = None,
        max_tokens: int = 3000,
    ) -> Tuple[List[str], List[Dict]]:
        """Retrieve the most relevant text chunks.

        Uses hybrid retrieval (semantic + BM25) followed by
        cross-encoder reranking.  Results below *confidence_threshold*
        are dropped so the LLM can abstain rather than hallucinate.

        Args:
            query: The search query string.
            top_k: Number of top results to return.
            confidence_threshold: Minimum reranker confidence (0-1).
                Results below this value are discarded.
            source_filter: If set, only return chunks whose
                ``source`` metadata matches (case-insensitive
                substring).
            max_tokens: Approximate token budget for total context
                sent to the LLM.  Chunks are added in score order
                until this budget is reached.

        Returns:
            Tuple of (documents, metadatas) with confidence.
        """
        initial_k = top_k * 3
        query_emb = self.embedding_model.encode(
            [query]
        ).tolist()

        # --- Stage 1a: Semantic (vector) search ---
        if self.backend == "faiss":
            results = self._faiss_query(
                query_emb[0], initial_k
            )
        else:
            results = self.collection.query(
                query_embeddings=query_emb,
                n_results=initial_k,
                include=["documents", "metadatas", "distances"],
            )

        if (
            not results
            or not results["documents"]
            or not results["documents"][0]
        ):
            return [], []

        flat_docs = results["documents"][0]
        flat_metas = (
            results["metadatas"][0]
            if results.get("metadatas")
            else [{}] * len(flat_docs)
        )
        flat_dists = (
            results["distances"][0]
            if results.get("distances")
            else [0.0] * len(flat_docs)
        )

        # --- Stage 1b: BM25 keyword search (hybrid) ---
        bm25_docs, bm25_metas = self._bm25_search(
            query, top_n=initial_k
        )

        # Merge BM25 results into the semantic pool (deduplicate)
        seen_texts = set(d[:200] for d in flat_docs)
        for doc, meta in zip(bm25_docs, bm25_metas):
            key = doc[:200]
            if key not in seen_texts:
                flat_docs.append(doc)
                flat_metas.append(meta)
                flat_dists.append(2.0)  # placeholder distance
                seen_texts.add(key)

        # --- Source filter ---
        if source_filter:
            sf_lower = source_filter.lower()
            filtered = [
                (d, m, dist)
                for d, m, dist in zip(flat_docs, flat_metas, flat_dists)
                if sf_lower in m.get("source", "").lower()
            ]
            if filtered:
                flat_docs, flat_metas, flat_dists = zip(*filtered)
                flat_docs = list(flat_docs)
                flat_metas = list(flat_metas)
                flat_dists = list(flat_dists)
            else:
                # Filter matched nothing → return empty (don't bypass filter)
                return [], []

        if not flat_docs:
            return [], []

        # --- Stage 2: Cross-encoder reranking ---
        pairs = [[query, doc] for doc in flat_docs]
        scores = self.reranker.predict(pairs)

        scored = self._score_results(
            flat_docs, flat_metas, scores, flat_dists
        )
        scored.sort(key=lambda x: x["score"], reverse=True)

        # --- Confidence threshold ---
        scored = [
            s for s in scored
            if s["meta"].get("confidence", 0) >= confidence_threshold
        ]

        # --- Token budget ---
        budget_remaining = max_tokens
        budget_results: List[Dict] = []
        for item in scored:
            est_tokens = len(item["doc"]) // 4
            if est_tokens > budget_remaining and budget_results:
                break
            budget_remaining -= est_tokens
            budget_results.append(item)
            if len(budget_results) >= top_k:
                break

        return (
            [item["doc"] for item in budget_results],
            [item["meta"] for item in budget_results],
        )

    # ------------------------------------------------------------------
    # BM25 keyword search
    # ------------------------------------------------------------------

    def _bm25_search(
        self, query: str, top_n: int = 10
    ) -> Tuple[List[str], List[Dict]]:
        """Run BM25 keyword search over all stored documents.

        Uses a cache to avoid rebuilding the BM25 index on every query.
        Cache is invalidated when corpus size changes.

        Args:
            query: Raw query string.
            top_n: Number of top results to return.

        Returns:
            Tuple of (documents, metadatas).
        """
        if not BM25_AVAILABLE:
            return [], []

        # Gather corpus from the active backend
        if self.backend == "faiss":
            corpus = self.faiss_documents
            metas = self.faiss_metadatas
            corpus_size = len(corpus)
        else:
            # Pull all documents from ChromaDB
            total = self.collection.count()
            if total == 0:
                return [], []
            corpus_size = total
            # Check if cache is valid (same corpus size)
            if (
                self._bm25_cache is not None
                and self._bm25_cache.get("corpus_size") == corpus_size
            ):
                corpus = self._bm25_cache["corpus"]
                metas = self._bm25_cache["metas"]
                bm25 = self._bm25_cache["bm25"]
            else:
                # Cache miss - rebuild BM25 index
                all_data = self.collection.get(
                    include=["documents", "metadatas"],
                    limit=total,
                )
                corpus = all_data.get("documents", [])
                metas = all_data.get("metadatas", []) or [{}] * len(corpus)
                if not corpus:
                    return [], []
                tokenized_corpus = [
                    doc.lower().split() for doc in corpus
                ]
                bm25 = BM25Okapi(tokenized_corpus)
                # Store in cache
                self._bm25_cache = {
                    "corpus_size": corpus_size,
                    "corpus": corpus,
                    "metas": metas,
                    "bm25": bm25,
                }

        if not corpus:
            return [], []

        # For FAISS backend, also use cache logic
        if self.backend == "faiss":
            if (
                self._bm25_cache is not None
                and self._bm25_cache.get("corpus_size") == corpus_size
            ):
                bm25 = self._bm25_cache["bm25"]
            else:
                tokenized_corpus = [
                    doc.lower().split() for doc in corpus
                ]
                bm25 = BM25Okapi(tokenized_corpus)
                self._bm25_cache = {
                    "corpus_size": corpus_size,
                    "corpus": corpus,
                    "metas": metas,
                    "bm25": bm25,
                }

        query_tokens = query.lower().split()
        scores = bm25.get_scores(query_tokens)

        top_indices = np.argsort(scores)[::-1][:top_n]
        docs_out = [corpus[i] for i in top_indices if scores[i] > 0]
        metas_out = [metas[i] for i in top_indices if scores[i] > 0]
        return docs_out, metas_out

    # ------------------------------------------------------------------
    # Document management
    # ------------------------------------------------------------------

    def list_documents(self) -> List[Dict]:
        """Return a list of unique source documents in the store.

        Returns:
            List of dicts with ``source``, ``chunk_count``, and
            ``total_chars`` keys.
        """
        if self.backend == "faiss":
            source_info: Dict[str, Dict] = {}
            for doc, meta in zip(
                self.faiss_documents, self.faiss_metadatas
            ):
                src = meta.get("source", "unknown")
                if src not in source_info:
                    source_info[src] = {
                        "source": src,
                        "chunk_count": 0,
                        "total_chars": 0,
                    }
                source_info[src]["chunk_count"] += 1
                source_info[src]["total_chars"] += len(doc)
            return list(source_info.values())

        # ChromaDB
        total = self.collection.count()
        if total == 0:
            return []
        all_data = self.collection.get(
            include=["documents", "metadatas"], limit=total
        )
        docs = all_data.get("documents", [])
        meta_list = all_data.get("metadatas", []) or [{}] * len(docs)
        source_info = {}
        for doc, meta in zip(docs, meta_list):
            src = meta.get("source", "unknown")
            if src not in source_info:
                source_info[src] = {
                    "source": src,
                    "chunk_count": 0,
                    "total_chars": 0,
                }
            source_info[src]["chunk_count"] += 1
            source_info[src]["total_chars"] += len(doc)
        return list(source_info.values())

    def delete_document(self, source_name: str) -> str:
        """Delete all chunks belonging to a specific source document.

        Args:
            source_name: The ``source`` metadata value
                (typically the filename).

        Returns:
            Status message.
        """
        if self.backend == "faiss":
            indices_to_keep = [
                i
                for i, m in enumerate(self.faiss_metadatas)
                if m.get("source", "") != source_name
            ]
            removed = len(self.faiss_documents) - len(indices_to_keep)
            if removed == 0:
                return f"No chunks found for '{source_name}'"

            new_docs = [self.faiss_documents[i] for i in indices_to_keep]
            new_metas = [self.faiss_metadatas[i] for i in indices_to_keep]
            new_ids = [self.faiss_ids[i] for i in indices_to_keep]

            # Rebuild FAISS index
            self.faiss_index = faiss.IndexFlatL2(self.embedding_dim)
            if new_docs:
                embeddings = self.embedding_model.encode(new_docs).tolist()
                vectors = np.array(embeddings, dtype=np.float32)
                self.faiss_index.add(vectors)

            self.faiss_documents = new_docs
            self.faiss_metadatas = new_metas
            self.faiss_ids = new_ids
            self._save_faiss()
            # Invalidate BM25 cache (corpus changed)
            self._bm25_cache = None
            return f"Deleted {removed} chunks for '{source_name}'"

        # ChromaDB
        total = self.collection.count()
        if total == 0:
            return f"No chunks found for '{source_name}'"

        all_data = self.collection.get(
            include=["metadatas"], limit=total
        )
        ids_to_delete = [
            doc_id
            for doc_id, meta in zip(
                all_data["ids"], all_data["metadatas"]
            )
            if meta.get("source", "") == source_name
        ]

        if not ids_to_delete:
            return f"No chunks found for '{source_name}'"

        # ChromaDB delete in batches (max 5461 per call)
        for i in range(0, len(ids_to_delete), 5000):
            batch = ids_to_delete[i : i + 5000]
            self.collection.delete(ids=batch)

        # Invalidate BM25 cache (corpus changed)
        self._bm25_cache = None

        return f"Deleted {len(ids_to_delete)} chunks for '{source_name}'"

    @staticmethod
    def _score_results(
        docs: List[str],
        metas: List[Dict],
        scores,
        distances: List[float] = None,
    ) -> List[Dict]:
        """Combine documents, metadata, and reranker scores.

        Args:
            docs: Retrieved document texts.
            metas: Retrieved metadata dicts.
            scores: Raw reranker scores.
            distances: L2 distances from ChromaDB vector search.

        Returns:
            List of scored result dicts.
        """
        if distances is None:
            distances = [0.0] * len(docs)
        scored: List[Dict] = []
        for doc, meta, score, dist in zip(
            docs, metas, scores, distances
        ):
            confidence = 1 / (1 + math.exp(-score))
            # ChromaDB default metric is L2 (squared).
            # Convert to cosine similarity: cos_sim = 1 - (L2^2 / 2)
            # (valid for normalised embeddings like all-MiniLM-L6-v2)
            cosine_sim = max(0.0, 1.0 - (dist / 2.0))
            euclidean = math.sqrt(max(0.0, dist))
            enriched = dict(meta) if meta else {}
            enriched["confidence"] = round(confidence, 3)
            enriched["rerank_score"] = round(float(score), 4)
            enriched["cosine_similarity"] = round(cosine_sim, 4)
            enriched["euclidean_distance"] = round(euclidean, 4)
            scored.append(
                {
                    "doc": doc,
                    "meta": enriched,
                    "score": score,
                }
            )
        return scored

    def analyze_query(self, query: str) -> dict:
        """Analyze query for search terms, rewrite, and source filter.

        Detects patterns like "what does <document> say about Y?"
        to enable multi-document querying awareness.

        Args:
            query: The user query string.

        Returns:
            Dict with query_rewrite, search_terms,
            retrieval_strategy, and optional source_filter keys.
        """
        stopwords = {
            "what", "is", "the", "a", "an", "how", "does",
            "do", "can", "could", "would", "should", "will",
            "are", "was", "were", "been", "being", "have",
            "has", "had", "of", "to", "for", "in", "on",
            "with", "by", "about", "this", "that", "these",
            "those", "it", "its", "and", "or",
        }

        words = re.findall(r"\b\w+\b", query.lower())
        search_terms = [
            w
            for w in words
            if w not in stopwords and len(w) > 2
        ]

        query_rewrite = query.strip()
        if query_rewrite.endswith("?"):
            query_rewrite = query_rewrite[:-1]

        # --- Multi-document awareness: detect source filter ---
        source_filter: Optional[str] = None
        # Patterns: "in document X", "from file X", "in X.pdf"
        source_patterns = [
            r"(?:in|from|according to)\s+(?:document|file|doc)\s+[\"']?(.+?)[\"']?\s*(?:,|\?|$)",
            r"(?:in|from)\s+[\"'](.+?\.\w{2,5})[\"']",
            r"(?:document|file)\s+[\"']?(\S+\.\w{2,5})[\"']?",
        ]
        for pattern in source_patterns:
            match = re.search(pattern, query, re.IGNORECASE)
            if match:
                source_filter = match.group(1).strip()
                break

        bm25_label = " + BM25 keyword" if BM25_AVAILABLE else ""
        strategy = (
            f"Hybrid search: semantic embeddings"
            f" (all-MiniLM-L6-v2){bm25_label}"
            f" → cross-encoder reranking"
        )
        result: dict = {
            "query_rewrite": query_rewrite,
            "search_terms": search_terms[:6],
            "retrieval_strategy": strategy,
        }
        if source_filter:
            result["source_filter"] = source_filter
        return result